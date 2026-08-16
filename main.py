# main.py

import logging
import re
import math
import zipfile
from xml.etree import ElementTree as ET
from typing import Optional
from pydantic import BaseModel
from pathlib import Path
from datetime import datetime
import asyncio
import shutil
import uuid
import razorpay

from fastapi import (
    FastAPI,
    UploadFile,
    File,
    Form,
    HTTPException,
    Request,
    Header
)

from vendor_routes import (
    router as vendor_router,
    get_authenticated_vendor
)

from fastapi.responses import JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.concurrency import run_in_threadpool

from pypdf import PdfReader
from pptx import Presentation
from docx import Document
from docx.oxml.ns import qn
import openpyxl

from vendor_database import (
    get_vendor_by_id,
    get_vendor_settings,
    update_vendor_settings,
    update_vendor_maintenance,
    update_vendor_accept_orders,
    get_vendor_maintenance,
)

from database import (
    initialize_database,
    save_print_job,
    get_print_job,
    get_all_print_jobs,
    update_payment_status,
    update_printer_status,
    update_job_details,
    get_next_queue_number,
    mark_payment_pending,
    mark_payment_success,
    mark_payment_failed,
    save_razorpay_order,
    get_razorpay_order_id,
    save_razorpay_payment,
    assign_queue_after_payment,
    start_printing,
    complete_printing,
    cleanup_expired_jobs,
    update_print_job,
    get_vendor_dashboard,
    get_vendor_orders,
    get_vendor_queue,
)


# ============================================================
# FASTAPI APP
# ============================================================

app = FastAPI(
    title="ServePrint API",
    description="Backend API for ServePrint",
    version="1.0.0"
)

app.include_router(vendor_router)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

logger = logging.getLogger(__name__)


# ============================================================
# FOLDERS
# ============================================================

UPLOAD_FOLDER = Path("uploads")
UPLOAD_FOLDER.mkdir(exist_ok=True)

LOG_FOLDER = Path("logs")
LOG_FOLDER.mkdir(exist_ok=True)


def cleanup_expired_uploads():
    expired_files = cleanup_expired_jobs()

    for filename in expired_files:
        file_path = UPLOAD_FOLDER / filename

        if file_path.exists():
            file_path.unlink()


# ============================================================
# BACKGROUND CLEANUP
# ============================================================

async def cleanup_scheduler():

    while True:

        try:
            cleanup_expired_uploads()
            logger.info("Expired jobs cleaned.")

        except Exception as e:
            logger.error(
                f"Cleanup Scheduler Error: {e}"
            )

        await asyncio.sleep(30)


initialize_database()


@app.on_event("startup")
async def startup_event():

    asyncio.create_task(
        cleanup_scheduler()
    )


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# STATIC UPLOADS
# ============================================================

app.mount(
    "/uploads",
    StaticFiles(
        directory="uploads"
    ),
    name="uploads"
)


# ============================================================
# FILE TYPES
# ============================================================

ALLOWED_EXTENSIONS = {
    ".pdf",

    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".heic",
    ".heif",

    ".docx",
    ".pptx",
    ".xlsx",
    ".txt"
}


ALLOWED_MIME_TYPES = {
    "application/pdf",

    "image/jpeg",
    "image/png",
    "image/gif",
    "image/webp",
    "image/bmp",
    "image/tiff",
    "image/heic",
    "image/heif",

    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",

    "application/vnd.openxmlformats-officedocument.presentationml.presentation",

    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

    "text/plain",

    "application/octet-stream"
}


# ============================================================
# PRICING
# ============================================================

BW_PRICE_PER_PAGE = 2.0
COLOR_PRICE_PER_PAGE = 10.0

MAX_FILE_SIZE = 20 * 1024 * 1024


def calculate_price(
    pages: int,
    copies: int = 1,
    print_type: str = "bw"
):

    rate = (
        COLOR_PRICE_PER_PAGE
        if print_type == "color"
        else BW_PRICE_PER_PAGE
    )

    return pages * copies * rate


def estimate_wait_time(
    queue_number: int
):
    return queue_number * 2


# ============================================================
# DOCX PAGE COUNTER
# ============================================================

def count_docx_pages(
    filepath: Path
) -> int:

    app_xml_pages = 0

    try:

        with zipfile.ZipFile(
            filepath
        ) as docx_zip:

            with docx_zip.open(
                "docProps/app.xml"
            ) as f:

                tree = ET.parse(f)

                ns = {
                    "ep":
                    "http://schemas.openxmlformats.org/"
                    "officeDocument/2006/extended-properties"
                }

                pages_el = tree.find(
                    "ep:Pages",
                    ns
                )

                if (
                    pages_el is not None
                    and
                    pages_el.text
                    and
                    pages_el.text.isdigit()
                ):

                    app_xml_pages = int(
                        pages_el.text
                    )

    except Exception as e:

        logger.warning(
            f"DOCX app.xml page count unavailable: {e}"
        )


    page_break_pages = 0
    word_count_pages = 0

    try:

        document = Document(
            str(filepath)
        )

        page_breaks = 0
        word_count = 0

        for paragraph in document.paragraphs:

            word_count += len(
                paragraph.text.split()
            )

            for run in paragraph.runs:

                for br in run._element.findall(
                    qn("w:br")
                ):

                    if (
                        br.get(
                            qn("w:type")
                        )
                        ==
                        "page"
                    ):

                        page_breaks += 1


        if page_breaks > 0:

            page_break_pages = (
                page_breaks + 1
            )


        if word_count > 0:

            word_count_pages = max(
                1,
                round(
                    word_count / 400
                )
            )

    except Exception as e:

        logger.error(
            f"DOCX fallback page count failed: {e}"
        )


    candidates = [
        p
        for p in (
            app_xml_pages,
            page_break_pages,
            word_count_pages
        )
        if p > 0
    ]


    return (
        max(candidates)
        if candidates
        else 1
    )


# ============================================================
# XLSX PAGE COUNTER
# ============================================================

def count_xlsx_pages(
    filepath: Path
) -> int:

    try:

        workbook = openpyxl.load_workbook(
            str(filepath),
            read_only=True
        )

        sheet_count = len(
            workbook.sheetnames
        )

        workbook.close()

        return max(
            1,
            sheet_count
        )

    except Exception as e:

        logger.error(
            f"XLSX page count failed: {e}"
        )

        return 1


# ============================================================
# TXT PAGE COUNTER
# ============================================================

def count_txt_pages(
    filepath: Path
) -> int:

    try:

        text = filepath.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        chars = len(text)

        if chars == 0:
            return 1

        return max(
            1,
            math.ceil(
                chars / 3000
            )
        )

    except Exception as e:

        logger.error(
            f"TXT page count failed: {e}"
        )

        return 1


# ============================================================
# UNIVERSAL PAGE COUNTER
# ============================================================

def count_pages(
    filepath: Path
):

    ext = filepath.suffix.lower()


    # ----------------------------
    # PDF
    # ----------------------------

    if ext == ".pdf":

        try:

            reader = PdfReader(
                str(filepath),
                strict=False
            )

            if reader.is_encrypted:

                try:
                    reader.decrypt("")

                except Exception:
                    pass


            pages = len(
                reader.pages
            )


            if pages < 1:

                raise ValueError(
                    "PDF contains no pages."
                )


            logger.info(
                f"PDF Pages: {pages}"
            )

            return pages


        except Exception as e:

            logger.exception(
                f"PDF ERROR: {e}"
            )

            raise ValueError(
                "We could not read this PDF. "
                "Please try a different PDF or "
                "re-save the file."
            ) from e


    # ----------------------------
    # PPTX
    # ----------------------------

    if ext == ".pptx":

        try:

            presentation = Presentation(
                str(filepath)
            )

            slides = len(
                presentation.slides
            )

            logger.info(
                f"PPTX Slides: {slides}"
            )

            return max(
                1,
                slides
            )

        except Exception as e:

            logger.error(
                f"PPTX ERROR: {e}"
            )

            return 1


    # ----------------------------
    # DOCX
    # ----------------------------

    if ext == ".docx":

        pages = count_docx_pages(
            filepath
        )

        logger.info(
            f"DOCX Pages: {pages}"
        )

        return pages


    # ----------------------------
    # XLSX
    # ----------------------------

    if ext == ".xlsx":

        pages = count_xlsx_pages(
            filepath
        )

        logger.info(
            f"XLSX Pages: {pages}"
        )

        return pages


    # ----------------------------
    # TXT
    # ----------------------------

    if ext == ".txt":

        pages = count_txt_pages(
            filepath
        )

        logger.info(
            f"TXT Pages: {pages}"
        )

        return pages


    # ----------------------------
    # Images
    # ----------------------------

    if ext in {
        ".jpg",
        ".jpeg",
        ".png",
        ".gif",
        ".webp",
        ".bmp",
        ".tif",
        ".tiff",
        ".heic",
        ".heif"
    }:

        return 1


    return 1


# ============================================================
# BASIC APIs
# ============================================================

@app.get("/")
def home():

    return {
        "project": "ServePrint",
        "status": "Running",
        "version": "1.0.0"
    }


@app.get("/health")
def health():

    return {
        "server": "Online",
        "printer": "Waiting",
        "database": "Connected"
    }


@app.get("/status")
def status():

    return {
        "printer": "Offline",
        "queue": 0,
        "jobs": 0
    }


# ============================================================
# PUBLIC VENDOR STATUS
#
# Customer receives ONLY:
# maintenance
# accept_orders
#
# Vendor ID is NOT returned.
# ============================================================

@app.get(
    "/vendor/{vendor_id}/status"
)
def vendor_status(
    vendor_id: str
):

    vendor = get_vendor_by_id(
        vendor_id
    )

    if vendor is None:

        raise HTTPException(
            status_code=404,
            detail="Vendor not found."
        )


    maintenance = bool(
        get_vendor_maintenance(
            vendor_id
        )
    )


    settings = get_vendor_settings(
        vendor_id
    )


    accept_orders = True


    if settings is not None:

        accept_orders = bool(
            settings.get(
                "accept_orders",
                True
            )
        )


    return {
        "maintenance":
            maintenance,

        "accept_orders":
            accept_orders
    }


# ============================================================
# UPLOAD API
# ============================================================

@app.post("/upload")
async def upload_file(

    vendor_id: str = Form(...),

    file: UploadFile = File(...),

    copies: int = Form(1),

    print_type: str = Form("bw"),

    paper_size: str = Form("A4"),

    page_range: str = Form("All")
):

    # ----------------------------
    # Vendor
    # ----------------------------

    vendor = get_vendor_by_id(
        vendor_id
    )

    if vendor is None:

        raise HTTPException(
            status_code=404,
            detail="Vendor not found."
        )


    # ----------------------------
    # Maintenance
    # ----------------------------

    maintenance = get_vendor_maintenance(
        vendor_id
    )

    if maintenance:

        raise HTTPException(
            status_code=503,
            detail=(
                "This shop is currently "
                "under maintenance. "
                "Please try again later."
            )
        )


    # ----------------------------
    # Accept Orders
    # ----------------------------

    settings = get_vendor_settings(
        vendor_id
    )

    if settings is not None:

        accept_orders = bool(
            settings.get(
                "accept_orders",
                True
            )
        )

        if not accept_orders:

            raise HTTPException(
                status_code=403,
                detail=(
                    "This shop is currently "
                    "not accepting new orders."
                )
            )


    # ----------------------------
    # Copies
    # ----------------------------

    if copies < 1:

        raise HTTPException(
            status_code=400,
            detail="Copies must be at least 1."
        )


    if copies > 100:

        raise HTTPException(
            status_code=400,
            detail="Maximum 100 copies allowed."
        )


    # ----------------------------
    # Page Range
    # ----------------------------

    page_range = (
        page_range.strip()
    )


    if page_range == "":

        raise HTTPException(
            status_code=400,
            detail="Page range cannot be empty."
        )


    if page_range.lower() != "all":

        pattern = r"^[0-9,\-\s]+$"

        if not re.match(
            pattern,
            page_range
        ):

            raise HTTPException(
                status_code=400,
                detail="Invalid page range."
            )


    # ----------------------------
    # File Name
    # ----------------------------

    if not file.filename:

        raise HTTPException(
            status_code=400,
            detail="No file selected."
        )


    extension = Path(
        file.filename
    ).suffix.lower()


    if extension not in ALLOWED_EXTENSIONS:

        raise HTTPException(
            status_code=400,
            detail="Unsupported file type."
        )


    # ----------------------------
    # Video Protection
    # ----------------------------

    content_type = (
        file.content_type or ""
    ).lower().strip()


    if content_type.startswith(
        "video/"
    ):

        raise HTTPException(
            status_code=400,
            detail="Video files are not allowed."
        )


    # ----------------------------
    # MIME
    # ----------------------------

    if (
        content_type
        and
        content_type not in ALLOWED_MIME_TYPES
        and
        not content_type.startswith(
            "image/"
        )
        and
        content_type !=
        "application/octet-stream"
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Please upload a PDF, image, "
                "DOCX, PPTX, XLSX or TXT file."
            )
        )


    # ----------------------------
    # Print Type
    # ----------------------------

    if print_type not in [
        "bw",
        "color"
    ]:

        raise HTTPException(
            status_code=400,
            detail="Invalid print type."
        )


    # ----------------------------
    # Job
    # ----------------------------

    job_id = str(
        uuid.uuid4()
    )


    filename = (
        f"{job_id}{extension}"
    )


    filepath = (
        UPLOAD_FOLDER / filename
    )


    # ----------------------------
    # Save
    # ----------------------------

    try:

        with open(
            filepath,
            "wb"
        ) as buffer:

            shutil.copyfileobj(
                file.file,
                buffer
            )

    except Exception as e:

        logger.exception(
            "File saving failed."
        )

        if filepath.exists():
            filepath.unlink()

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to save the uploaded file."
            )
        ) from e


    # ----------------------------
    # Size
    # ----------------------------

    try:

        file_size = (
            filepath.stat().st_size
        )

    except Exception as e:

        if filepath.exists():
            filepath.unlink()

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to read uploaded file."
            )
        ) from e


    if file_size <= 0:

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )


    if file_size > MAX_FILE_SIZE:

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail="Maximum file size is 20 MB."
        )


    # ----------------------------
    # Count Pages
    # ----------------------------

    try:

        total_pages = (
            await run_in_threadpool(
                count_pages,
                filepath
            )
        )

    except ValueError as e:

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail=str(e)
        ) from e


    except Exception as e:

        logger.exception(
            "File processing failed."
        )

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "The server could not "
                "process this file. "
                "Please try another file."
            )
        ) from e


    if total_pages < 1:

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "The file contains "
                "no printable pages."
            )
        )


    logger.info(
        f"Detected Pages: {total_pages}"
    )


    # ----------------------------
    # Queue
    # ----------------------------

    queue_number = 0
    waiting_time = 0


    # ----------------------------
    # Price
    # ----------------------------

    total_amount = calculate_price(
        pages=total_pages,
        copies=copies,
        print_type=print_type
    )


    # ----------------------------
    # Database
    # ----------------------------

    try:

        save_print_job(
            vendor_id=vendor_id,
            job_id=job_id,
            original_name=file.filename,
            saved_name=filename,
            file_size=file_size,
            total_pages=total_pages,
            queue_number=queue_number,
            total_amount=total_amount
        )


        update_job_details(
            job_id,
            total_pages,
            total_amount,
            queue_number,
            copies,
            print_type,
            paper_size
        )


        mark_payment_pending(
            job_id
        )


    except Exception as e:

        logger.exception(
            "Database job creation failed."
        )

        filepath.unlink(
            missing_ok=True
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Unable to create the "
                "print job. Please try again."
            )
        ) from e


    # ----------------------------
    # Customer Response
    #
    # NO vendor_id
    # ----------------------------

    return {

        "success":
            True,

        "job_id":
            job_id,

        "original_name":
            file.filename,

        "saved_name":
            filename,

        "file_size":
            file_size,

        "uploaded_at":
            datetime.now().isoformat(),

        "queue_number":
            0,

        "total_pages":
            total_pages,

        "total_amount":
            total_amount,

        "estimated_wait_time":
            waiting_time,

        "copies":
            copies,

        "print_type":
            print_type,

        "paper_size":
            paper_size,

        "page_range":
            page_range,

        "payment_status":
            "Pending",

        "printer_status":
            "Pending Payment"
    }


# ============================================================
# PRINT JOB
# ============================================================

class PrintJobRequest(
    BaseModel
):

    job_id: str
    copies: int
    print_type: str
    paper_size: str

    orientation: Optional[str] = (
        "Portrait"
    )

    page_range: str


@app.post("/print")
def create_print_job(
    request: PrintJobRequest
):

    update_print_job(
        job_id=request.job_id,
        copies=request.copies,
        print_type=request.print_type,
        paper_size=request.paper_size,
        orientation=request.orientation,
        page_range=request.page_range
    )


    return {

        "success":
            True,

        "message":
            "Print Job Created",

        "job_id":
            request.job_id
    }

# ============================================================
# UPDATE EXISTING JOB
# ============================================================

@app.put(
    "/jobs/{job_id}"
)
def update_existing_job(

    job_id: str,

    request: PrintJobRequest
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found"
        )


    if job[
        "payment_status"
    ] == "Paid":

        raise HTTPException(
            status_code=400,
            detail=(
                "Paid jobs cannot be modified."
            )
        )


    total_pages = job[
        "total_pages"
    ]


    total_amount = calculate_price(
        pages=total_pages,
        copies=request.copies,
        print_type=request.print_type
    )


    update_print_job(
        job_id=job_id,
        copies=request.copies,
        print_type=request.print_type,
        paper_size=request.paper_size,
        orientation=request.orientation,
        page_range=request.page_range
    )


    update_job_details(
        job_id,
        total_pages,
        total_amount,
        job["queue_number"],
        request.copies,
        request.print_type,
        request.paper_size
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "total_pages":
            total_pages,

        "total_amount":
            total_amount,

        "queue_number":
            job["queue_number"],

        "estimated_wait_time":
            max(
                0,
                (
                    job["queue_number"]
                    - 1
                ) * 2
            ),

        "printer_status":
            job["printer_status"]
    }


# ============================================================
# GET SINGLE JOB
# ============================================================

@app.get(
    "/jobs/{job_id}"
)
def fetch_job(
    job_id: str
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job Not Found"
        )


    job = dict(job)


    job["orders_ahead"] = max(
        0,
        job["queue_number"] - 1
    )


    job[
        "estimated_wait_time"
    ] = (
        job["orders_ahead"] * 2
    )


    # Never expose vendor ID
    job.pop(
        "vendor_id",
        None
    )


    return job


# ============================================================
# GET ALL JOBS
# ============================================================

@app.get("/jobs")
def fetch_all_jobs():

    jobs = get_all_print_jobs()

    result = []

    for job in jobs:

        item = dict(job)

        item.pop(
            "vendor_id",
            None
        )

        result.append(
            item
        )

    return result


# ============================================================
# PAYMENT STATUS
# ============================================================

@app.put(
    "/jobs/{job_id}/payment/{status}"
)
def update_job_payment_status(
    job_id: str,
    status: str
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    allowed = [
        "Pending",
        "Paid",
        "Failed"
    ]


    if status not in allowed:

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid payment status."
            )
        )


    update_payment_status(
        job_id,
        status
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "payment_status":
            status
    }


# ============================================================
# PRINTER STATUS
# ============================================================

@app.put(
    "/jobs/{job_id}/printer/{status}"
)
def update_job_printer_status(
    job_id: str,
    status: str
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    allowed = [
        "Pending Payment",
        "Waiting",
        "Printing",
        "Completed",
        "Failed"
    ]


    if status not in allowed:

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid printer status."
            )
        )


    update_printer_status(
        job_id,
        status
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "printer_status":
            status
    }

# ============================================================
# CREATE RAZORPAY ORDER
# ============================================================

@app.post(
    "/payment/create/{job_id}"
)
def create_razorpay_order(
    job_id: str
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    if job[
        "payment_status"
    ] == "Paid":

        raise HTTPException(
            status_code=400,
            detail=(
                "This job has already been paid."
            )
        )


    if job[
        "payment_status"
    ] != "Pending":

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment cannot be created "
                "from "
                f"'{job['payment_status']}' "
                "state."
            )
        )


    # ----------------------------
    # Vendor Credentials
    # ----------------------------

    settings = get_vendor_settings(
        job["vendor_id"]
    )


    if settings is None:

        raise HTTPException(
            status_code=404,
            detail=(
                "Vendor settings not found."
            )
        )


    key_id = (
        settings.get(
            "razorpay_key"
        )
        or ""
    ).strip()


    key_secret = (
        settings.get(
            "razorpay_secret"
        )
        or ""
    ).strip()


    if not key_id or not key_secret:

        raise HTTPException(
            status_code=400,
            detail=(
                "Razorpay is not configured "
                "for this vendor."
            )
        )


    amount_rupees = float(
        job["total_amount"] or 0
    )


    amount_paise = int(
        round(
            amount_rupees * 100
        )
    )


    if amount_paise < 100:

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment amount must be "
                "at least ₹1."
            )
        )


    try:

        client = razorpay.Client(
            auth=(
                key_id,
                key_secret
            )
        )


        order = client.order.create({

            "amount":
                amount_paise,

            "currency":
                "INR",

            "receipt":
                job_id[:40],

            "notes": {
                "job_id":
                    job_id,

                "vendor_id":
                    job["vendor_id"]
            }
        })


    except Exception as e:

        logger.exception(
            "Razorpay order creation failed."
        )

        raise HTTPException(
            status_code=502,
            detail=(
                "Unable to create "
                "Razorpay order."
            )
        ) from e


    razorpay_order_id = (
        order.get("id")
    )


    if not razorpay_order_id:

        raise HTTPException(
            status_code=502,
            detail=(
                "Razorpay did not return "
                "an order ID."
            )
        )


    save_razorpay_order(
        job_id,
        razorpay_order_id
    )


    # Vendor ID remains internal.
    # Do not return it to customer.

    return {

        "success":
            True,

        "job_id":
            job_id,

        "razorpay_order_id":
            razorpay_order_id,

        "razorpay_key_id":
            key_id,

        "amount":
            amount_paise,

        "amount_rupees":
            amount_rupees,

        "currency":
            "INR",

        "payment_status":
            "Pending"
    }


# ============================================================
# RAZORPAY PAYMENT VERIFICATION
# ============================================================

class RazorpayPaymentVerification(
    BaseModel
):

    razorpay_payment_id: str

    razorpay_signature: str


@app.post(
    "/payment/verify/{job_id}"
)
def verify_razorpay_payment(

    job_id: str,

    data:
        RazorpayPaymentVerification
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    if job[
        "payment_status"
    ] == "Paid":

        return {

            "success":
                True,

            "job_id":
                job_id,

            "payment_status":
                "Paid",

            "queue_number":
                job["queue_number"],

            "message":
                "Payment already completed."
        }


    if job[
        "payment_status"
    ] != "Pending":

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment cannot be verified "
                "from "
                f"'{job['payment_status']}' "
                "state."
            )
        )


    razorpay_order_id = (
        get_razorpay_order_id(
            job_id
        )
    )


    if not razorpay_order_id:

        raise HTTPException(
            status_code=400,
            detail=(
                "Razorpay order was not "
                "created for this job."
            )
        )


    settings = get_vendor_settings(
        job["vendor_id"]
    )


    if settings is None:

        raise HTTPException(
            status_code=404,
            detail=(
                "Vendor settings not found."
            )
        )


    key_id = (
        settings.get(
            "razorpay_key"
        )
        or ""
    ).strip()


    key_secret = (
        settings.get(
            "razorpay_secret"
        )
        or ""
    ).strip()


    if not key_id or not key_secret:

        raise HTTPException(
            status_code=400,
            detail=(
                "Razorpay is not configured "
                "for this vendor."
            )
        )


    try:

        client = razorpay.Client(
            auth=(
                key_id,
                key_secret
            )
        )


        client.utility.verify_payment_signature({

            "razorpay_order_id":
                razorpay_order_id,

            "razorpay_payment_id":
                data.razorpay_payment_id,

            "razorpay_signature":
                data.razorpay_signature
        })


    except Exception:

        logger.exception(
            "Razorpay signature verification failed."
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid Razorpay payment signature."
            )
        )


    try:

        payment = client.payment.fetch(
            data.razorpay_payment_id
        )

    except Exception:

        logger.exception(
            "Unable to fetch Razorpay payment."
        )

        raise HTTPException(
            status_code=502,
            detail=(
                "Unable to verify payment "
                "status with Razorpay."
            )
        )


    razorpay_payment_order_id = (
        payment.get("order_id")
    )


    if (
        razorpay_payment_order_id
        != razorpay_order_id
    ):

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment does not belong "
                "to this order."
            )
        )


    expected_amount = int(
        round(
            float(
                job["total_amount"] or 0
            ) * 100
        )
    )


    received_amount = int(
        payment.get(
            "amount",
            0
        )
    )


    if (
        received_amount
        != expected_amount
    ):

        raise HTTPException(
            status_code=400,
            detail="Payment amount mismatch."
        )


    payment_status = (
        payment.get("status")
    )


    if payment_status != "captured":

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment is not captured. "
                f"Current status: "
                f"{payment_status}"
            )
        )


    save_razorpay_payment(
        job_id,
        data.razorpay_payment_id,
        data.razorpay_signature
    )


    try:

        queue_number = (
            assign_queue_after_payment(
                job["vendor_id"],
                job_id,
                data.razorpay_payment_id
            )
        )

    except ValueError as error:

        raise HTTPException(
            status_code=400,
            detail=str(error)
        )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "payment_id":
            data.razorpay_payment_id,

        "razorpay_order_id":
            razorpay_order_id,

        "payment_status":
            "Paid",

        "queue_number":
            queue_number,

        "estimated_wait_time":
            max(
                0,
                (
                    queue_number - 1
                ) * 2
            ),

        "printer_status":
            "Waiting"
    }


# ============================================================
# PAYMENT FAILED
# ============================================================

@app.post(
    "/payment/{job_id}/failed"
)
def payment_failed(
    job_id: str
):

    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    if job[
        "payment_status"
    ] == "Paid":

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment is already completed."
            )
        )


    if job[
        "payment_status"
    ] == "Failed":

        return {

            "success":
                True,

            "job_id":
                job_id,

            "payment_status":
                "Failed",

            "message":
                "Payment is already "
                "marked as failed."
        }


    if job[
        "payment_status"
    ] != "Pending":

        raise HTTPException(
            status_code=400,
            detail=(
                "Payment cannot be marked "
                "failed from "
                f"'{job['payment_status']}' "
                "state."
            )
        )


    mark_payment_failed(
        job_id
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "payment_status":
            "Failed",

        "queue_number":
            0,

        "printer_status":
            "Pending Payment",

        "message":
            "Payment failed. "
            "No queue number assigned."
    }

# ============================================================
# START PRINTING
# ============================================================

@app.post(
    "/jobs/{job_id}/start"
)
def start_job(

    job_id: str,

    authorization:
        str | None =
        Header(default=None)
):

    vendor = get_authenticated_vendor(
        authorization
    )


    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    if (
        job["vendor_id"]
        !=
        vendor["vendor_id"]
    ):

        raise HTTPException(
            status_code=403,
            detail=(
                "You are not authorized "
                "to start this job."
            )
        )


    start_printing(
        job_id
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "printer_status":
            "Printing"
    }


# ============================================================
# COMPLETE PRINTING
# ============================================================

@app.post(
    "/jobs/{job_id}/complete"
)
def complete_job(

    job_id: str,

    authorization:
        str | None =
        Header(default=None)
):

    vendor = get_authenticated_vendor(
        authorization
    )


    job = get_print_job(
        job_id
    )


    if not job:

        raise HTTPException(
            status_code=404,
            detail="Job not found."
        )


    if (
        job["vendor_id"]
        !=
        vendor["vendor_id"]
    ):

        raise HTTPException(
            status_code=403,
            detail=(
                "You are not authorized "
                "to complete this job."
            )
        )


    complete_printing(
        job_id
    )


    return {

        "success":
            True,

        "job_id":
            job_id,

        "printer_status":
            "Completed"
    }


# ============================================================
# VENDOR DASHBOARD
# ============================================================

@app.get(
    "/vendor/{vendor_id}/dashboard"
)
def vendor_dashboard(
    vendor_id: str
):

    return get_vendor_dashboard(
        vendor_id
    )


# ============================================================
# VENDOR ORDERS
# ============================================================

@app.get(
    "/vendor/{vendor_id}/orders"
)
def vendor_orders(
    vendor_id: str
):

    return get_vendor_orders(
        vendor_id
    )


# ============================================================
# VENDOR ORDERS
# ============================================================

@app.get(
    "/vendor/{vendor_id}/orders"
)
def vendor_orders(
    vendor_id: str
):

    return get_vendor_orders(
        vendor_id
    )


# ============================================================
# VENDOR QUEUE
# ============================================================

@app.get(
    "/vendor/{vendor_id}/queue"
)
def vendor_queue(
    vendor_id: str
):

    return get_vendor_queue(
        vendor_id
    )


# ============================================================
# VENDOR QR
# ============================================================

@app.get(
    "/vendor/{vendor_id}/qr"
)
def vendor_qr(
    vendor_id: str
):

    vendor = get_vendor_by_id(
        vendor_id
    )


    if vendor is None:

        raise HTTPException(
            status_code=404,
            detail="Vendor not found"
        )


    frontend_url = (
        "https://server-point-1vrst74ki-jyotirmaya01s-projects.vercel.app"
    )


    return {

        "vendor_id":
            vendor["vendor_id"],

        "shop_name":
            vendor["shop_name"],

        "url":
            (
                f"{frontend_url}/"
                f"?vendor_id="
                f"{vendor['vendor_id']}"
            )
    }


# ============================================================
# VENDOR SETTINGS MODEL
# ============================================================

class VendorSettingsUpdate(
    BaseModel
):

    shop_name: str

    owner_name: str

    phone: str

    address: str

    maintenance: bool

    accept_orders: bool

    razorpay_key: str = ""

    razorpay_secret: str = ""

    google_sheet_id: str = ""

    service_email: str = ""

    smtp_host: str = ""

    smtp_port: int = 587

    smtp_email: str = ""

    smtp_password: str = ""


# ============================================================
# VENDOR SETTINGS
# ============================================================

@app.get(
    "/vendor/{vendor_id}/settings"
)
def vendor_settings(
    vendor_id: str
):

    settings = get_vendor_settings(
        vendor_id
    )


    if settings is None:

        raise HTTPException(
            status_code=404,
            detail=(
                "Vendor settings not found."
            )
        )


    return settings


# ============================================================
# UPDATE VENDOR SETTINGS
# ============================================================

@app.put(
    "/vendor/{vendor_id}/settings"
)
def save_vendor_settings(

    vendor_id: str,

    data:
        VendorSettingsUpdate
):

    update_vendor_settings(
        vendor_id,
        data
    )


    return {

        "success":
            True,

        "message":
            "Settings updated successfully."
    }


# ============================================================
# VENDOR MAINTENANCE
# ============================================================

@app.post(
    "/vendor/{vendor_id}/maintenance"
)
def set_vendor_maintenance(

    vendor_id: str,

    enabled: bool
):

    vendor = get_vendor_by_id(
        vendor_id
    )


    if vendor is None:

        raise HTTPException(
            status_code=404,
            detail="Vendor not found."
        )


    success = (
        update_vendor_maintenance(
            vendor_id,
            enabled
        )
    )


    if not success:

        raise HTTPException(
            status_code=404,
            detail=(
                "Vendor settings not found."
            )
        )


    return {

        "success":
            True,

        "vendor_id":
            vendor_id,

        "maintenance":
            enabled
    }


# ============================================================
# VENDOR ACCEPT ORDERS
# ============================================================

@app.post(
    "/vendor/{vendor_id}/accept-orders"
)
def set_vendor_accept_orders(

    vendor_id: str,

    enabled: bool
):

    vendor = get_vendor_by_id(
        vendor_id
    )


    if vendor is None:

        raise HTTPException(
            status_code=404,
            detail="Vendor not found."
        )


    success = (
        update_vendor_accept_orders(
            vendor_id,
            enabled
        )
    )


    if not success:

        raise HTTPException(
            status_code=404,
            detail=(
                "Vendor settings not found."
            )
        )


    return {

        "success":
            True,

        "vendor_id":
            vendor_id,

        "accept_orders":
            enabled
    }


# ============================================================
# CUSTOMER VENDOR VALIDATION
#
# IMPORTANT:
# vendor_id is NEVER returned.
# ============================================================

@app.get(
    "/customer/vendor/{vendor_id}/validate"
)
def validate_customer_vendor(
    vendor_id: str
):

    vendor = get_vendor_by_id(
        vendor_id
    )


    if vendor is None:

        return {

            "valid":
                False,

            "message":
                "Shop not found."
        }


    maintenance = (
        get_vendor_maintenance(
            vendor_id
        )
    )


    if maintenance is None:
        maintenance = False


    settings = get_vendor_settings(
        vendor_id
    )


    if settings is None:

        accept_orders = True

    else:

        accept_orders = bool(
            settings.get(
                "accept_orders",
                True
            )
        )


    return {

        "valid":
            True,

        "maintenance":
            maintenance,

        "accept_orders":
            accept_orders,

        "active":
            True
    }


# ============================================================
# GLOBAL EXCEPTION HANDLER
# ============================================================

@app.exception_handler(
    Exception
)
async def global_exception_handler(
    request: Request,
    exc: Exception
):

    logger.error(
        f"SERVER ERROR: {exc}"
    )

    logger.exception(
        "Unhandled exception"
    )


    return JSONResponse(
        status_code=500,
        content={

            "success":
                False,

            "message":
                "Internal Server Error"
        }
    )